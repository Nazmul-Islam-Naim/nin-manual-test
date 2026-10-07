import io
import re
import uuid

from docx import Document
from pypdf import PdfWriter


def err(r, status, code):
    assert r.status_code == status, r.text
    body = r.json()
    assert set(body) == {"error"} and body["error"]["code"] == code
    assert isinstance(body["error"]["message"], str) and body["error"]["message"]


def upload(client, name, data):
    return client.post("/manuals", files={"file": (name, data)})


def make_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)
    return out


def test_paste_text_roundtrip(client):
    r = client.post("/manuals", json={"text": "  Step 1: খুলুন\r\nStep 2: Click\rEnd  "})
    assert r.status_code == 201
    b = r.json()
    assert uuid.UUID(b["id"]).version == 4
    assert b["source_type"] == "text" and b["original_filename"] is None
    assert "text" not in b and "file_path" not in b
    assert re.fullmatch(r"\d{4}-\d\d-\d\dT.*Z", b["created_at"])
    g = client.get(f"/manuals/{b['id']}")
    assert g.status_code == 200
    gb = g.json()
    assert gb["text"] == "Step 1: খুলুন\nStep 2: Click\nEnd"
    assert gb["char_count"] == len(gb["text"]) == b["char_count"]
    assert "file_path" not in gb


def test_txt_and_md_bangla(client):
    for name, kind in (("a.txt", "txt"), ("B.MD", "md")):
        r = upload(client, name, "বাংলা and English\r\nline2".encode("utf-8"))
        assert r.status_code == 201, r.text
        assert r.json()["source_type"] == kind and r.json()["original_filename"] == name
        assert client.get(f"/manuals/{r.json()['id']}").json()["text"] == "বাংলা and English\nline2"


def test_invalid_utf8(client):
    err(upload(client, "a.txt", b"\xff\xfe\x00bad\x80"), 422, "invalid_encoding")


def test_docx_paragraphs_and_tables(client):
    d = Document()
    d.add_paragraph("Hello বাংলা")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "cellA"
    t.cell(0, 1).text = "cellB"
    buf = io.BytesIO()
    d.save(buf)
    r = upload(client, "m.docx", buf.getvalue())
    assert r.status_code == 201, r.text
    text = client.get(f"/manuals/{r.json()['id']}").json()["text"]
    assert "Hello বাংলা" in text and "cellA" in text and "cellB" in text


def test_pdf_text_layer(client):
    r = upload(client, "a.pdf", make_pdf("Login steps"))
    assert r.status_code == 201, r.text
    assert r.json()["source_type"] == "pdf"
    assert "Login steps" in client.get(f"/manuals/{r.json()['id']}").json()["text"]


def test_pdf_without_text_layer(client):
    w = PdfWriter()
    w.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    w.write(buf)
    err(upload(client, "scan.pdf", buf.getvalue()), 422, "no_readable_text")


def test_corrupt_files(client):
    err(upload(client, "x.pdf", b"not a pdf"), 422, "invalid_request")
    err(upload(client, "x.docx", b"not a zip"), 422, "invalid_request")


def test_empty_inputs(client):
    err(client.post("/manuals", json={"text": "  \n\t "}), 422, "empty_text")
    err(client.post("/manuals", json={"text": ""}), 422, "empty_text")
    err(upload(client, "e.txt", b"   \n"), 422, "no_readable_text")


def test_unsupported_type(client):
    err(upload(client, "a.exe", b"x"), 415, "unsupported_file_type")
    err(upload(client, "noext", b"x"), 415, "unsupported_file_type")


def test_too_large(client):
    err(upload(client, "big.txt", b"a" * (1024 * 1024 + 1)), 413, "file_too_large")
    assert upload(client, "ok.txt", b"a" * (1024 * 1024)).status_code == 201


def test_invalid_requests(client):
    err(client.post("/manuals", json={}), 422, "invalid_request")
    err(client.post("/manuals", json={"text": 5}), 422, "invalid_request")
    err(client.post("/manuals", content=b"{bad", headers={"content-type": "application/json"}), 422, "invalid_request")
    err(client.post("/manuals", content=b"x", headers={"content-type": "text/plain"}), 422, "invalid_request")
    err(client.post("/manuals", data={"other": "x"}, files={"zzz": ("a.txt", b"x")}), 422, "invalid_request")


def test_file_wins_over_text(client):
    r = client.post("/manuals", data={"text": "pasted"}, files={"file": ("a.txt", b"from file")})
    assert r.status_code == 201 and r.json()["source_type"] == "txt"


def test_not_found(client):
    err(client.get(f"/manuals/{uuid.uuid4()}"), 404, "manual_not_found")


def test_unknown_route_uses_error_shape(client):
    r = client.get("/nope")
    assert r.status_code == 404 and "error" in r.json()
